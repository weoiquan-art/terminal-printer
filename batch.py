"""M3：批量打印。先生成 Excel 表（按货号带出 DESC/MODEL），改好后再打印。

用法（Windows）：
    python batch.py prepare "81 96020 0340" "WG 9525160001" -o out\\batch.xlsx --open
    python batch.py prepare --from codes.txt -o out\\batch.xlsx      每行一个货号
    （在 Excel 里核对 DESC/MODEL，填 PRICE 和 QTY，保存）
    python batch.py print out\\batch.xlsx                          只检查并出总览图
    python batch.py print out\\batch.xlsx --print                  检查通过后打印

print 会先检查整张表，任何一行有问题都不打印。
"""

import argparse
import csv
import os
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

from labelkit import layout, preview, printer, tspl
from labelkit.parts import Parts

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_TEMPLATE = os.path.join(HERE, "templates", "default_50x40.json")
COLUMNS = ["NO", "DESC", "MODEL", "PRICE", "QTY", "NOTE"]
WIDTHS = [20, 34, 46, 12, 6, 40]


def cmd_prepare(args):
    codes = list(args.codes)
    if args.from_file:
        with open(args.from_file, encoding="utf-8-sig") as f:
            codes += [line.strip() for line in f if line.strip()]
    if not codes:
        sys.exit("错误：没有货号")

    parts = Parts()
    wb = Workbook()
    ws = wb.active
    ws.title = "labels"
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    found = 0
    for code in codes:
        hit, desc, model, note = parts.lookup(code)
        found += hit is not None
        ws.append([hit or code, desc, model, "", None, note])
    for i, width in enumerate(WIDTHS):
        col = ws.cell(row=1, column=i + 1).column_letter
        ws.column_dimensions[col].width = width
    # 货号、文字列设为文本格式，防止 Excel 把 81960200340 改成 8.2E+10
    for row in ws.iter_rows(min_row=2, max_col=4):
        for cell in row:
            cell.number_format = "@"
            cell.alignment = Alignment(wrap_text=False)
    ws.freeze_panes = "A2"

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    wb.save(args.out)
    print("已生成 {}：{} 个货号，{} 个在数据里找到".format(args.out, len(codes), found))
    print("请核对 DESC/MODEL，填 PRICE 和 QTY，保存后运行：python batch.py print {}".format(args.out))
    if args.open:
        os.startfile(os.path.abspath(args.out))


def _read_rows(path):
    """返回 [(Excel 行号, {列名: 值}), ...]，跳过全空行。"""
    if path.lower().endswith(".csv"):
        with open(path, encoding="utf-8-sig", newline="") as f:
            reader = csv.reader(f)
            header = [h.strip().upper() for h in next(reader)]
            raw = list(enumerate(reader, 2))
    else:
        ws = load_workbook(path, data_only=True).active
        it = ws.iter_rows(values_only=True)
        header = [str(h or "").strip().upper() for h in next(it)]
        raw = list(enumerate(it, 2))
    missing = [c for c in COLUMNS[:5] if c not in header]
    if missing:
        sys.exit("错误：表里缺少列 {}".format(", ".join(missing)))
    rows = []
    for n, values in raw:
        rec = dict(zip(header, values))
        if all(v is None or str(v).strip() == "" for k, v in rec.items() if k != "NOTE"):
            continue
        rows.append((n, rec))
    return rows


def _text(v):
    return "" if v is None else str(v).strip()


def cmd_print(args):
    tpl = layout.load_template(args.template)
    try:
        parts = Parts()
    except FileNotFoundError:
        parts = None
    rows = _read_rows(args.path)
    if not rows:
        sys.exit("错误：表里没有要打印的行")

    errors, warnings, jobs, items = [], [], [], []
    label = None
    for n, rec in rows:
        no_raw = rec.get("NO")
        if isinstance(no_raw, (int, float)):
            errors.append("第 {} 行：NO 是数字 {!r}，可能被 Excel 改过，请把该列设为文本后重填".format(n, no_raw))
            continue
        no = _text(no_raw)
        qty_raw = rec.get("QTY")
        try:
            qty = int(float(_text(qty_raw)))
            if qty < 1 or qty != float(_text(qty_raw)):
                raise ValueError
        except ValueError:
            errors.append("第 {} 行（{}）：QTY 要填正整数，现在是 {!r}".format(n, no, qty_raw))
            continue
        values = {"no": no, "desc": _text(rec.get("DESC")),
                  "model": _text(rec.get("MODEL")), "price": _text(rec.get("PRICE"))}
        for key in ("no", "desc", "price"):
            if not values[key]:
                errors.append("第 {} 行（{}）：{} 不能为空".format(n, no, key.upper()))
        if any(not values[k] for k in ("no", "desc", "price")):
            continue
        if parts and parts.lookup(no)[0] != no:
            warnings.append("第 {} 行：货号 {} 不在零件数据里（照样打印，请确认没打错）".format(n, no))
        try:
            label, ops = layout.build(tpl, values)
        except ValueError as e:
            errors.append("第 {} 行（{}）：{}".format(n, no, e))
            continue
        jobs.append((ops, qty))
        items.append(("row {}  {}  x{}".format(n, no, qty), ops))

    for w in warnings:
        print("提醒：" + w)
    if errors:
        print("\n".join("错误：" + e for e in errors))
        sys.exit("共 {} 处错误，没有打印。改好表格后再运行。".format(len(errors)))

    total = sum(q for _, q in jobs)
    out = args.preview or os.path.splitext(os.path.abspath(args.path))[0] + "_preview.png"
    preview.sheet(label, items).save(out)
    print("检查通过：{} 种标签，共 {} 张。总览图：{}".format(len(jobs), total, out))

    data = tspl.render_many(label, jobs)
    if args.dry_run:
        sys.stdout.write(tspl.render_many(label, jobs, display=True))
    if args.do_print:
        written = printer.send_raw(args.printer, data, "terminal-printer batch")
        print("已发送 {} 字节到「{}」，共 {} 张".format(written, args.printer, total))
    else:
        print("未打印。确认总览图无误后加 --print。")


def main():
    for stream in (sys.stdout, sys.stderr):
        stream.reconfigure(encoding="utf-8", errors="replace")
    parser = argparse.ArgumentParser(description="批量标签")
    sub = parser.add_subparsers(dest="cmd", required=True)

    p = sub.add_parser("prepare", help="按货号生成 Excel 表")
    p.add_argument("codes", nargs="*", help="货号（含空格时加引号）")
    p.add_argument("--from", dest="from_file", help="货号列表文件，每行一个")
    p.add_argument("-o", "--out", default=os.path.join("out", "batch.xlsx"))
    p.add_argument("--open", action="store_true", help="生成后用 Excel 打开")
    p.set_defaults(func=cmd_prepare)

    p = sub.add_parser("print", help="检查表格、出总览图，可选打印")
    p.add_argument("path", help="batch.xlsx 或 .csv")
    p.add_argument("--template", default=DEFAULT_TEMPLATE)
    p.add_argument("--preview", help="总览图路径，默认与表格同名 _preview.png")
    p.add_argument("--dry-run", action="store_true", help="输出 TSPL")
    p.add_argument("--print", dest="do_print", action="store_true", help="发送到打印机")
    p.add_argument("--printer", default=printer.DEFAULT_PRINTER)
    p.set_defaults(func=cmd_print)

    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
